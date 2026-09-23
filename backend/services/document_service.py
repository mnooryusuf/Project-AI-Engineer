"""
services/document_service.py — Memproses dokumen untuk RAG pipeline
Mendukung: PDF, TXT, DOCX, XLSX
"""
import os
import aiofiles
from pathlib import Path
from typing import List
from sqlalchemy.orm import Session
from models import Document
from services.embedding_service import get_embedding
from config import get_settings

settings = get_settings()

# Ukuran chunk. Dibatasi 300 karakter karena paraphrase-multilingual hanya
# menerima 128 token per input: chunk 500 karakter TERBUKTI ditolak Ollama
# dengan HTTP 500 pada 11 dari 50 chunk saat diuji. Yang gagal semuanya
# berasal dari PDF hasil scan — ekstraksi pypdf menyisipkan spasi di antara
# hampir tiap huruf ("D o k u m e n"), sehingga jumlah token meledak jauh di
# atas perkiraan dari jumlah karakter. Teks .txt bersih dengan panjang sama
# lolos tanpa masalah. 300 karakter aman untuk kedua jenis teks.
CHUNK_SIZE = 300
CHUNK_OVERLAP = 50

# Halaman PDF dengan teks di bawah ambang ini dianggap hasil pindaian dan
# di-OCR. 50 karakter cukup rendah untuk tidak salah menuduh halaman yang
# memang isinya sedikit (mis. halaman lampiran berisi satu baris), tapi cukup
# tinggi untuk menangkap halaman scan yang biasanya menghasilkan 0 karakter
# atau beberapa karakter sampah dari artefak PDF.
MIN_PAGE_TEXT_CHARS = 50

# Batas halaman yang di-OCR per dokumen. OCR berjalan ~33 detik per halaman
# di mesin target (lihat PDF_OCR_DPI di tools/ocr_tool.py), jadi tanpa batas
# ini sebuah PDF pindaian 30 halaman akan menahan permintaan upload selama
# belasan menit. 10 halaman (~5,5 menit) sudah mencakup hampir semua surat
# dinas; sisanya sengaja dilewati dan dilaporkan ke pengguna, bukan
# dikerjakan diam-diam sampai koneksi putus.
MAX_OCR_PAGES = 10

# Batas yang sama saat Apple Vision aktif (tools/ocr_tool.uses_fast_ocr):
# ~1 detik per halaman, jadi 60 halaman masih sekitar satu menit.
MAX_OCR_PAGES_FAST = 60


def _page_has_images(page) -> bool:
    """Halaman PDF memuat gambar (XObject /Image) — dicek dari resource
    halaman tanpa mendekode gambarnya, supaya murah untuk PDF besar."""
    try:
        xobjects = page.get("/Resources", {}).get("/XObject", {})
        return any(obj.get_object().get("/Subtype") == "/Image" for obj in xobjects.values())
    except Exception:
        return False


def _normalize(text: str) -> str:
    return "".join(ch for ch in text.lower() if ch.isalnum())


def _new_ocr_lines(ocr_text: str, text_layer: str) -> list[str]:
    """Baris OCR yang BELUM ada di lapisan teks halaman. Teks ketikan asli
    selalu lebih akurat, jadi yang ditambahkan hanya isi gambarnya saja."""
    layer = _normalize(text_layer)
    return [
        line for line in ocr_text.splitlines()
        if len(_normalize(line)) >= 3 and _normalize(line) not in layer
    ]


def chunk_text(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> List[str]:
    """Membagi teks menjadi chunk-chunk kecil dengan overlap."""
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunks.append(text[start:end])
        start += chunk_size - overlap
    return chunks


async def extract_text_from_file(file_path: str) -> str:
    """Ekstrak teks dari file PDF, TXT, DOCX, atau XLSX."""
    ext = Path(file_path).suffix.lower()

    if ext == ".txt":
        async with aiofiles.open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return await f.read()

    elif ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        halaman = [(page.extract_text() or "") for page in reader.pages]

        # PDF hasil pindaian isinya gambar, bukan teks — pypdf mengembalikan
        # string kosong tanpa error apa pun, sehingga dokumen diam-diam masuk
        # knowledge base dalam keadaan kosong. Halaman semacam itu di-OCR.
        # Dicek PER HALAMAN, bukan per dokumen, karena surat dinas sering
        # bercampur: halaman ketikan digital ditambah halaman tanda tangan
        # hasil scan. Halaman yang sudah punya teks tidak di-OCR ulang —
        # teks aslinya selalu lebih akurat daripada hasil pembacaan gambar.
        from tools.ocr_tool import extract_text_from_pdf_pages, uses_fast_ocr

        perlu_ocr = [i for i, t in enumerate(halaman) if len(t.strip()) < MIN_PAGE_TEXT_CHARS]

        # Halaman yang SUDAH berteks tapi memuat gambar juga dibaca, karena
        # isi pentingnya sering ada di gambar: di berita acara survei harga,
        # lapisan teks hanya berisi nama barang, sedangkan harganya
        # (Rp17.355.000, Rp10.960.000, ...) ada di tangkapan layar
        # marketplace yang ditempel — tidak pernah terbaca sebelumnya. Hanya
        # dengan Apple Vision; dengan EasyOCR tiap halaman ~33 detik.
        fast = uses_fast_ocr()
        tambahan = []
        if fast:
            tambahan = [
                i for i, page in enumerate(reader.pages)
                if i not in perlu_ocr and _page_has_images(page)
            ]

        target = (perlu_ocr + tambahan)[: MAX_OCR_PAGES_FAST if fast else MAX_OCR_PAGES]
        if target:
            hasil_ocr = await extract_text_from_pdf_pages(file_path, target)
            for i, teks in hasil_ocr.items():
                if i in perlu_ocr:
                    halaman[i] = teks
                else:
                    baru = _new_ocr_lines(teks, halaman[i])
                    if baru:
                        halaman[i] += "\n[Teks dalam gambar]\n" + "\n".join(baru)

        return "\n".join(halaman)

    elif ext == ".docx":
        # Alias "DocxFile" — modul ini juga mengimpor model ORM bernama
        # Document (baris import di atas), jadi python-docx.Document tidak
        # boleh dipakai dengan nama yang sama.
        from docx import Document as DocxFile
        docx_file = DocxFile(file_path)
        parts = [p.text for p in docx_file.paragraphs if p.text.strip()]
        # Tabel tidak ikut kebaca lewat .paragraphs (API python-docx
        # memisahkan keduanya) — banyak surat/KAK dinas menaruh info penting
        # di tabel (nomor, tanggal, rincian anggaran), jadi ikut diekstrak.
        for table in docx_file.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts)

    elif ext == ".xlsx":
        from openpyxl import load_workbook
        # read_only=True: baca streaming tanpa memuat seluruh workbook ke
        # RAM sekaligus — penting untuk mesin 8GB yang jadi target project ini.
        workbook = load_workbook(file_path, data_only=True, read_only=True)
        parts = []
        for sheet in workbook.worksheets:
            parts.append(f"# Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    parts.append(" | ".join(cells))
        workbook.close()
        return "\n".join(parts)

    else:
        raise ValueError(f"Format file tidak didukung: {ext}")


async def store_text_as_document(
    text: str,
    filename: str,
    db: Session,
    extra_metadata: dict | None = None,
) -> int:
    """
    Chunk + embedding + simpan teks ke knowledge base.

    Dipisah dari process_and_store_document() karena tidak semua sumber teks
    berasal dari file yang bisa dibaca extract_text_from_file() — teks hasil
    OCR gambar masuk lewat jalur ini juga (lihat /upload di main.py).

    Mengembalikan jumlah chunk yang disimpan; 0 berarti teksnya kosong dan
    TIDAK ada apa pun yang tersimpan — pemanggil wajib memperlakukan itu
    sebagai kegagalan, bukan keberhasilan.
    """
    chunks = chunk_text(text)
    saved = 0

    for i, chunk in enumerate(chunks):
        if not chunk.strip():
            continue
        embedding = await get_embedding(chunk)
        doc = Document(
            filename=filename,
            content=chunk,
            embedding=embedding,
            doc_metadata={"chunk_index": i, "total_chunks": len(chunks), **(extra_metadata or {})},
        )
        db.add(doc)
        saved += 1

    db.commit()
    return saved


async def process_and_store_document(
    file_path: str,
    filename: str,
    db: Session
) -> int:
    """
    Pipeline lengkap: ekstrak teks dari file -> chunk -> embedding -> simpan.
    Mengembalikan jumlah chunk yang disimpan.
    """
    text = await extract_text_from_file(file_path)
    return await store_text_as_document(text, filename, db)
